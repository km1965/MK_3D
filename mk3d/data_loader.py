"""Chargement et validation des fichiers de points XYZ (Excel / CSV)."""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".csv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}

X_ALIASES = ("x", "x (m)", "x(m)", "east", "est", "easting", "abscisse")
Y_ALIASES = ("y", "y (m)", "y(m)", "north", "nord", "northing", "ordonnee", "ordonnée")
Z_ALIASES = ("z", "z (m)", "z(m)", "alt", "altitude", "elev", "elevation", "élévation", "cote", "côte")

LABEL_ALIASES = (
    "libelle", "libellé", "label", "nom", "name",
    "designation", "désignation", "point", "point_id", "id_point",
    "numero", "n°", "num", "batiment", "bâtiment", "immeuble", "ouvrage",
    "groupe", "group", "code", "repere", "repère",
)

# Mots-clés indiquant le niveau bas d'un profil (pied du talus).
LEVEL_LOW_KEYWORDS = (
    "bas", "basse", "pied", "inf", "inférieur", "inferieure", "inferieur",
    "aval", "bottom",
)
# Mots-clés indiquant le niveau haut d'un profil (crête du talus).
LEVEL_HIGH_KEYWORDS = (
    "haut", "haute", "crête", "crete", "sommet", "top", "sup",
    "supérieur", "superieure", "superieur", "amont", "dessus",
)


def strip_accents(text: str) -> str:
    """Retire les accents d'une chaîne (ex. 'crête' -> 'crete')."""
    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(c for c in decomposed if not unicodedata.combining(c))


def normalise_label(text: object) -> str:
    """Normalise un libellé : minuscules, sans accents, séparateurs unifiés."""
    raw = strip_accents(str(text)).strip().lower()
    for sep in ("_", "-", ".", "/", "\\", "|", ",", ";"):
        raw = raw.replace(sep, " ")
    return " ".join(raw.split())


def split_label_level(text: object) -> tuple[str, str | None]:
    """Sépare un libellé en (côté, niveau).

    Le niveau vaut ``"bas"`` ou ``"haut"`` lorsqu'un mot-clé de niveau est
    présent, sinon ``None``. Le côté est le libellé privé du mot de niveau.

    Exemples
    --------
    >>> split_label_level("RD Talus bas")
    ('rd talus', 'bas')
    >>> split_label_level("Crête RG")
    ('rg', 'haut')
    """
    words = normalise_label(text).split()
    if not words:
        return "", None
    level: str | None = None
    side_words: list[str] = []
    for word in words:
        if word in LEVEL_LOW_KEYWORDS:
            level = "bas"
        elif word in LEVEL_HIGH_KEYWORDS:
            level = "haut"
        else:
            side_words.append(word)
    return " ".join(side_words), level


class LoadError(Exception):
    """Erreur de chargement d'un fichier de points."""


@dataclass
class XYZData:
    frame: pd.DataFrame
    x_col: str
    y_col: str
    z_col: str
    label_col: str | None = None
    source: Path | None = None
    dropped: int = 0

    @property
    def size(self) -> int:
        return len(self.frame)

    @property
    def x(self) -> np.ndarray:
        return self.frame[self.x_col].to_numpy(dtype=float)

    @property
    def y(self) -> np.ndarray:
        return self.frame[self.y_col].to_numpy(dtype=float)

    @property
    def z(self) -> np.ndarray:
        return self.frame[self.z_col].to_numpy(dtype=float)

    @property
    def has_labels(self) -> bool:
        return self.label_col is not None and self.label_col in self.frame.columns

    @property
    def labels(self) -> pd.Series:
        if self.has_labels:
            return self.frame[self.label_col].astype(str)
        return pd.Series([""] * len(self.frame))

    def statistics(self) -> pd.DataFrame:
        rows = []
        for label, col, values in (
            ("X", self.x_col, self.x),
            ("Y", self.y_col, self.y),
            ("Z", self.z_col, self.z),
        ):
            rows.append(
                {
                    "Axe": label,
                    "Colonne": col,
                    "Minimum": float(np.min(values)),
                    "Maximum": float(np.max(values)),
                    "Moyenne": float(np.mean(values)),
                    "Ecart-type": float(np.std(values)),
                }
            )
        return pd.DataFrame(rows)


def _normalise(name: object) -> str:
    return str(name).strip().lower().replace("_", " ")


def _find_column(columns, aliases) -> str | None:
    normalised = {_normalise(c): c for c in columns}
    for alias in aliases:
        if alias in normalised:
            return normalised[alias]
    return None


def _detect_columns(df: pd.DataFrame) -> tuple[str, str, str]:
    x_col = _find_column(df.columns, X_ALIASES)
    y_col = _find_column(df.columns, Y_ALIASES)
    z_col = _find_column(df.columns, Z_ALIASES)

    if x_col and y_col and z_col:
        return x_col, y_col, z_col

    numeric = list(df.select_dtypes(include="number").columns)
    if len(numeric) >= 3:
        return numeric[0], numeric[1], numeric[2]

    remaining = [c for c in df.columns if c not in {x_col, y_col, z_col}]
    candidates = numeric + remaining
    if len(candidates) >= 3:
        return candidates[0], candidates[1], candidates[2]

    raise LoadError(
        "Impossible d'identifier 3 colonnes numériques (X, Y, Z). "
        "Vérifiez que le fichier contient bien des coordonnées numériques."
    )


def _read_raw(path: Path, sheet: str | int | None = None) -> pd.DataFrame:
    if path.suffix.lower() in EXCEL_SUFFIXES:
        if path.suffix.lower() == ".xls":
            try:
                return pd.read_excel(path, sheet_name=sheet if sheet is not None else 0)
            except ImportError as exc:  # pragma: no cover
                raise LoadError(
                    "Le format .xls nécessite le paquet 'xlrd' "
                    "(pip install xlrd). Préférez un fichier .xlsx."
                ) from exc
        return pd.read_excel(path, sheet_name=sheet if sheet is not None else 0, engine="openpyxl")

    for sep in (None, ";", ",", "\t"):
        try:
            df = pd.read_csv(path, sep=sep, engine="python", decimal=".")
            if df.shape[1] >= 2:
                return df
        except Exception:
            continue
    raise LoadError("Lecture CSV impossible. Vérifiez le séparateur et l'encodage du fichier.")


def list_sheets(path: str | Path) -> list[str]:
    """Renvoie la liste des feuilles d'un classeur Excel (vide si non applicable)."""
    path = Path(path)
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        return []
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    except Exception:
        return []


def load_points(path: str | Path, sheet: str | int | None = None) -> XYZData:
    """Charge un fichier XYZ et renvoie un objet :class:`XYZData` validé."""
    path = Path(path)
    if not path.exists():
        raise LoadError(f"Fichier introuvable : {path}")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise LoadError(
            f"Format non supporté : '{path.suffix}'. "
            f"Formats acceptés : {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )

    try:
        raw = _read_raw(path, sheet)
    except LoadError:
        raise
    except Exception as exc:
        raise LoadError(f"Erreur de lecture du fichier : {exc}") from exc

    if raw.empty:
        raise LoadError("Le fichier ne contient aucune ligne de données.")

    x_col, y_col, z_col = _detect_columns(raw)

    label_col = _find_column(raw.columns, LABEL_ALIASES)

    cols = [x_col, y_col, z_col]
    if label_col and label_col not in cols:
        cols.append(label_col)

    frame = raw[cols].copy()
    rename_map = {x_col: "X", y_col: "Y", z_col: "Z"}
    if label_col:
        rename_map[label_col] = "LIBELLE"
    frame = frame.rename(columns=rename_map)

    for c in ["X", "Y", "Z"]:
        frame[c] = pd.to_numeric(frame[c], errors="coerce")
    before = len(frame)
    frame = frame.dropna(subset=["X", "Y", "Z"]).reset_index(drop=True)
    dropped = before - len(frame)
    if frame.empty:
        raise LoadError("Aucune ligne valide (X, Y, Z num�riques) n'a pu �tre extraite.")

    return XYZData(
        frame=frame,
        x_col="X",
        y_col="Y",
        z_col="Z",
        label_col="LIBELLE" if label_col and "LIBELLE" in frame.columns else None,
        source=path,
        dropped=dropped,
    )
