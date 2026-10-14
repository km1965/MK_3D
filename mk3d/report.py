"""Génération du rapport Word (.docx) contenant les figures produites."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .data_loader import XYZData
from .plotting import FigureResult

ACCENT = RGBColor(0x1F, 0x4E, 0x79)


@dataclass
class ReportOptions:
    title: str = "Rapport d'analyse XYZ"
    author: str = ""
    subject: str = "Visualisation 3D de données de points"
    include_table: bool = True


def _add_title_page(doc: Document, data: XYZData, opts: ReportOptions) -> None:
    title = doc.add_heading(opts.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(opts.subject)
    run.italic = True
    run.font.size = Pt(14)
    run.font.color.rgb = ACCENT

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Source : {data.source.name}\n"
        f"Nombre de points : {data.size}\n"
        f"Date de génération : {datetime.now():%d/%m/%Y %H:%M}"
    )

    if opts.author:
        author_p = doc.add_paragraph()
        author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_p.add_run(f"Auteur : {opts.author}").italic = True

    doc.add_page_break()


def _shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _add_statistics(doc: Document, data: XYZData) -> None:
    doc.add_heading("1. Statistiques descriptives", level=1)
    doc.add_paragraph(
        "Le tableau ci-dessous synthétise les statistiques des coordonnées "
        "des points importés (valeurs arrondies au millimètre)."
    )

    stats = data.statistics()
    table = doc.add_table(rows=1, cols=len(stats.columns))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header = table.rows[0].cells
    for idx, col in enumerate(stats.columns):
        header[idx].text = str(col)
        for paragraph in header[idx].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
        _shade_cell(header[idx], "1F4E79")

    for _, row in stats.iterrows():
        cells = table.add_row().cells
        for idx, col in enumerate(stats.columns):
            value = row[col]
            cells[idx].text = f"{value:.3f}" if isinstance(value, float) else str(value)

    if getattr(data, "dropped", 0):
        note = doc.add_paragraph()
        run = note.add_run(
            f"Remarque : {data.dropped} ligne(s) invalide(s) ont été ignorées lors de l'import."
        )
        run.italic = True
        run.font.size = Pt(9)


def _add_figures(doc: Document, figures: list[FigureResult]) -> None:
    doc.add_heading("2. Figures et visualisations", level=1)
    doc.add_paragraph(
        "Les figures ci-après présentent le nuage de points tridimensionnel "
        "ainsi que les projections orthogonales dans les plans X-Y, X-Z et Y-Z."
    )
    for figure in figures:
        if not figure.path.exists():
            continue
        doc.add_heading(figure.title, level=2)
        doc.add_picture(str(figure.path), width=Inches(6.1))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run(figure.caption)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def build_report(
    data: XYZData,
    figures: list[FigureResult],
    out_path: str | Path,
    opts: ReportOptions | None = None,
) -> Path:
    """Construit et enregistre le rapport Word. Renvoie le chemin du fichier."""
    opts = opts or ReportOptions()
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    _add_title_page(doc, data, opts)

    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(
        f"Ce rapport présente l'analyse géométrique d'un ensemble de {data.size} points "
        f"issus du fichier « {data.source.name} ». Les coordonnées ont été traitées à "
        "l'aide de l'application MK_3D afin de produire des figures 3D et des projections "
        "2D exploitables pour l'interprétation des données."
    )

    if opts.include_table:
        _add_statistics(doc, data)

    _add_figures(doc, figures)

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph(
        "L'ensemble des figures ci-dessus permet d'appréhender la répartition "
        "spatiale des points analysés et la variation de l'altitude Z sur l'emprise étudiée."
    )

    footer = doc.sections[0].footer
    footer.paragraphs[0].text = f"MK_3D - {datetime.now():%d/%m/%Y}"
    footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(out_path))
    return out_path
